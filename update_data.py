#!/usr/bin/env python3
"""Обновление данных рынка для калькулятора окупаемости.

Что делает:
1. Спрашивает у Банка России ключевую ставку из трёх независимых источников и сверяет их между собой.
2. Читает базу программ банков из baza-bankov.xlsx (лист «Программы»).
3. Пишет всё в data/market.json, который читает сайт.

Про ключевую ставку важно различать три даты:
  effective_since — когда ставку изменили в последний раз. Если ЦБ сохраняет ставку, она не меняется
                    месяцами, и это нормально;
  confirmed_at    — когда мы в последний раз подтвердили значение по данным ЦБ. По ней судят о свежести;
  last_meeting    — итог последнего заседания (только если его удалось прочитать в этом запуске).

Запуск:            python update_data.py
Проверка свежести: python update_data.py --check   (код возврата 1, если ставка не подтверждена свежо)
"""
import argparse
import datetime as dt
import json
import re
import sys
import time
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
from pathlib import Path

CBR_SOAP_URL = "https://www.cbr.ru/DailyInfoWebServ/DailyInfo.asmx"
CBR_SOAP_ACTION = '"http://web.cbr.ru/KeyRateXML"'
CBR_TABLE_URL = "https://www.cbr.ru/hd_base/KeyRate/"
CBR_PRESS_URL = "https://www.cbr.ru/press/keypr/"

STALE_DAYS = 14              # строка базы банков старше этого срока считается устаревшей
KEY_RATE_MAX_AGE_DAYS = 3    # ставка, не подтверждённая дольше этого срока, считается устаревшей
KEY_RATE_WINDOW_DAYS = 400   # окно запроса к ЦБ: ставку могут не менять больше года
MEETING_MAX_AGE_DAYS = 100   # итог заседания старше этого срока на сайте не показываем
SHEET = "Программы"
FIRST_ROW = 3

COL = {
    "bank": 2, "name": 3, "object_type": 4, "borrower": 5, "market": 6,
    "rate_min": 7, "rate_max": 8, "rate_conditions": 9, "down_min": 10,
    "term_max_years": 11, "amount_min": 12, "amount_max": 13,
    "rental_allowed": 14, "insurance_required": 15, "fees": 16,
    "source_url": 17, "checked": 18, "notes": 20,
}

MONTHS = {"января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6, "июля": 7,
          "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12}
DECISIONS = {"сохранить": "сохранена", "снизить": "снижена", "повысить": "повышена"}


# ---------- Банк России: разбор ответов ----------

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _collect_pairs(root):
    pairs = []
    for el in root.iter():
        kids = {_local(k.tag): (k.text or "").strip() for k in el}
        if "DT" in kids and "Rate" in kids:
            try:
                pairs.append((dt.date.fromisoformat(kids["DT"][:10]),
                              float(kids["Rate"].replace(",", "."))))
            except ValueError:
                continue
    return pairs


def _latest_run(pairs):
    """Пары (дата, ставка) -> (дата, ставка), где дата это начало последнего отрезка с той же ставкой.

    ЦБ может отдавать ставку на каждый день, а не только на дни изменения, поэтому «самая поздняя
    дата в ответе» не значит «ставку изменили в этот день». Если отрезок тянется до самого начала
    запрошенного окна, настоящее начало неизвестно: возвращаем None, и дата берётся из прошлого запуска.
    """
    pairs = sorted(set(pairs))
    i = len(pairs) - 1
    rate = pairs[i][1]
    while i > 0 and abs(pairs[i - 1][1] - rate) < 1e-9:
        i -= 1
    if i == 0 and len(pairs) > 1:
        return None, rate
    return pairs[i][0], rate


def parse_key_rate(xml_text: str):
    """Ответ веб-сервиса ЦБ (KeyRateXML) -> (дата начала действия текущей ставки, ставка)."""
    root = ET.fromstring(xml_text)
    pairs = _collect_pairs(root)
    if not pairs:  # результат мог прийти текстом с экранированным XML
        for el in root.iter():
            if _local(el.tag).endswith("Result") and el.text and "<" in el.text:
                pairs = _collect_pairs(ET.fromstring(el.text))
                break
    if not pairs:
        raise ValueError("в ответе ЦБ не нашлось пар «дата, ставка»")
    return _latest_run(pairs)


def _strip_tags(html: str) -> str:
    text = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
    text = re.sub(r"<[^>]+>", " ", text)
    text = text.replace("&nbsp;", " ").replace("\xa0", " ")
    return re.sub(r"\s+", " ", text)


def parse_key_rate_table(html: str, today: dt.date):
    """Страница ЦБ с таблицей «дата, ставка» -> последняя пара (дата, ставка)."""
    text = _strip_tags(html)
    pairs = []
    for m in re.finditer(r"(\d{2})\.(\d{2})\.(\d{4})\s+(\d{1,2}[.,]\d{1,2})(?!\d)", text):
        try:
            d = dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            rate = float(m.group(4).replace(",", "."))
        except ValueError:
            continue
        if d <= today and 0 < rate <= 50:
            pairs.append((d, rate))
    if not pairs:
        raise ValueError("в таблице ЦБ не нашлось пар «дата, ставка»")
    return _latest_run(pairs)


_PRESS_RE = re.compile(
    r"(сохранить|снизить|повысить)\s+ключевую\s+ставку\s+"
    r"(?:на\s+уровне\s+|(?:на\s+[\d.,]+\s*б\.?\s*п\.?,?\s*)?до\s+)"
    r"(\d{1,2}[.,]\d{1,2})\s*%", re.I)
_DATE_RE = re.compile(r"(\d{1,2})\s+(" + "|".join(MONTHS) + r")\s+(\d{4})|(\d{2})\.(\d{2})\.(\d{4})", re.I)


def parse_press(html: str):
    """Страница пресс-релизов ЦБ о ключевой ставке -> (ставка, итог заседания или None)."""
    text = _strip_tags(html)
    m = _PRESS_RE.search(text)
    if not m:
        raise ValueError("на странице ЦБ не нашлось решения по ключевой ставке")
    rate = float(m.group(2).replace(",", "."))
    verb = DECISIONS[m.group(1).lower()]
    dates = []
    for d in _DATE_RE.finditer(text):
        try:
            if d.group(1):
                val = dt.date(int(d.group(3)), MONTHS[d.group(2).lower()], int(d.group(1)))
            else:
                val = dt.date(int(d.group(6)), int(d.group(5)), int(d.group(4)))
        except ValueError:
            continue
        dates.append((d.start(), val))
    before = [v for pos, v in dates if pos < m.start()]
    date = before[-1] if before else (dates[0][1] if dates else None)
    return rate, ({"date": date, "decision": verb} if date else None)


# ---------- Банк России: запросы ----------

def _http(url: str, data=None, headers=None, timeout: int = 30, attempts: int = 3) -> str:
    req = urllib.request.Request(url, data=data, method="POST" if data else "GET",
                                 headers=dict({"User-Agent": "realty-data-updater/2.0"}, **(headers or {})))
    last = None
    for i in range(attempts):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                return resp.read().decode("utf-8", errors="replace")
        except (urllib.error.URLError, TimeoutError, OSError) as e:
            last = e
            time.sleep(2 * (i + 1))
    raise RuntimeError(f"{url}: {last}")


def _soap_body(date_from: dt.date, date_to: dt.date) -> bytes:
    return (
        '<?xml version="1.0" encoding="utf-8"?>'
        '<soap:Envelope xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" '
        'xmlns:xsd="http://www.w3.org/2001/XMLSchema" '
        'xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/">'
        '<soap:Body><KeyRateXML xmlns="http://web.cbr.ru/">'
        f"<fromDate>{date_from.isoformat()}</fromDate>"
        f"<ToDate>{date_to.isoformat()}</ToDate>"
        "</KeyRateXML></soap:Body></soap:Envelope>"
    ).encode("utf-8")


def src_soap(today: dt.date) -> dict:
    body = _soap_body(today - dt.timedelta(days=KEY_RATE_WINDOW_DAYS), today)
    text = _http(CBR_SOAP_URL, data=body, headers={
        "Content-Type": "text/xml; charset=utf-8", "SOAPAction": CBR_SOAP_ACTION})
    d, rate = parse_key_rate(text)
    return {"source": "soap", "value_pct": rate, "effective_since": d, "meeting": None}


def src_table(today: dt.date) -> dict:
    start = today - dt.timedelta(days=KEY_RATE_WINDOW_DAYS)
    url = (f"{CBR_TABLE_URL}?UniDbQuery.Posted=True&UniDbQuery.From={start:%d.%m.%Y}"
           f"&UniDbQuery.To={today:%d.%m.%Y}")
    d, rate = parse_key_rate_table(_http(url), today)
    return {"source": "table", "value_pct": rate, "effective_since": d, "meeting": None}


def src_press(today: dt.date) -> dict:
    rate, meeting = parse_press(_http(CBR_PRESS_URL))
    if meeting and not (0 <= (today - meeting["date"]).days <= MEETING_MAX_AGE_DAYS):
        meeting = None
    return {"source": "press", "value_pct": rate, "effective_since": None, "meeting": meeting}


# порядок важен: первый успешный источник с датой изменения становится основным
SOURCES = [("soap", src_soap), ("table", src_table), ("press", src_press)]


def collect_key_rate(today: dt.date):
    results, errors = [], []
    for name, fn in SOURCES:
        try:
            results.append(fn(today))
        except (RuntimeError, ValueError, ET.ParseError) as e:
            errors.append(f"источник «{name}»: {e}")
    return results, errors


def _iso(d):
    return d.isoformat() if d else None


def _age_days(iso, today):
    try:
        return (today - dt.date.fromisoformat(iso)).days
    except (TypeError, ValueError):
        return None


def reconcile_key_rate(results, errors, previous, today):
    """Сводит ответы источников в один блок key_rate. Возвращает (блок, предупреждения)."""
    warnings = ["ЦБ: " + e for e in errors]
    prev = previous or {}

    if results:
        values = {round(r["value_pct"], 3) for r in results}
        primary = next((r for r in results if r["effective_since"]), results[0])
        value = primary["value_pct"]
        effective = primary["effective_since"]
        if effective is None and prev.get("value_pct") == value:
            effective = _parse_date(prev.get("effective_since"))
        meeting = next((r["meeting"] for r in results if r["meeting"]), None)
        conflict = len(values) > 1
        if conflict:
            warnings.append("источники ЦБ дали разные значения ставки: " + ", ".join(
                f"{r['source']}={r['value_pct']}" for r in results))
        block = {
            "value_pct": value,
            "effective_since": _iso(effective),
            "confirmed_at": today.isoformat(),
            "last_meeting": ({"date": _iso(meeting["date"]), "decision": meeting["decision"]} if meeting else None),
            "sources": [r["source"] for r in results],
        }
        status = "conflict" if conflict else "ok"
    else:
        block = {
            "value_pct": prev.get("value_pct"),
            "effective_since": prev.get("effective_since"),
            "confirmed_at": prev.get("confirmed_at"),
            "last_meeting": None,
            "sources": [],
        }
        status = None
        if errors:
            warnings.append("ни один источник ЦБ не ответил, ставка взята из прошлого запуска")

    block["status"] = _status(block, today) if status is None else status
    return block, warnings


def _parse_date(iso):
    try:
        return dt.date.fromisoformat(iso) if iso else None
    except ValueError:
        return None


def _status(kr, today):
    if kr.get("status") == "conflict":
        return "conflict"
    if kr.get("value_pct") is None:
        return "missing"
    age = _age_days(kr.get("confirmed_at"), today)
    return "ok" if age is not None and age <= KEY_RATE_MAX_AGE_DAYS else "stale"


# ---------- База программ ----------

def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def _date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def load_programs(xlsx_path: Path, today: dt.date):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET]
    programs, warnings = [], []
    for r in range(FIRST_ROW, ws.max_row + 1):
        g = lambda k: ws.cell(r, COL[k]).value
        if not g("bank"):
            continue
        label = f"строка {r} ({g('bank')}, {g('name')})"
        checked = _date(g("checked"))
        age = (today - checked).days if checked else None
        pct = lambda v: round(v * 100, 3) if _num(v) is not None else None
        p = {
            "bank": g("bank"), "name": g("name"), "object_type": g("object_type"),
            "borrower": g("borrower"), "market": g("market"),
            "rate_min_pct": pct(g("rate_min")), "rate_max_pct": pct(g("rate_max")),
            "rate_conditions": g("rate_conditions"),
            "down_min_pct": pct(g("down_min")), "term_max_years": _num(g("term_max_years")),
            "amount_min": _num(g("amount_min")), "amount_max": _num(g("amount_max")),
            "rental_allowed": g("rental_allowed") or "Уточнить",
            "insurance_required": g("insurance_required"), "fees": g("fees"),
            "source_url": g("source_url"),
            "checked": checked.isoformat() if checked else None,
            "age_days": age, "stale": age is None or age > STALE_DAYS,
            "notes": g("notes"),
        }
        if p["rate_min_pct"] is None:
            warnings.append(f"{label}: не указана минимальная ставка")
        if not p["source_url"]:
            warnings.append(f"{label}: нет ссылки на источник")
        if checked is None:
            warnings.append(f"{label}: нет даты проверки")
        elif checked > today:
            warnings.append(f"{label}: дата проверки в будущем")
        if p["rate_min_pct"] is not None and not (0 < p["rate_min_pct"] < 80):
            warnings.append(f"{label}: ставка {p['rate_min_pct']}% выглядит нереалистично")
        programs.append(p)
    return programs, warnings


# ---------- Сборка и проверка ----------

def apply_verification(programs, verification, today):
    """Учесть ежедневную сверку с источниками (data/verification.json).

    ok      -> дата проверки = день сверки (но не раньше даты в базе и не из будущего);
    changed -> программа считается устаревшей, на сайте показывается ставка с источника;
    прочее  -> ничего не меняется.
    """
    sources = (verification or {}).get("sources") or {}
    for p in programs:
        v = sources.get(p.get("source_url") or "")
        p["verification"] = None
        if not v:
            continue
        p["verification"] = v.get("status")
        day = _parse_date(v.get("checked_at"))
        if v.get("status") == "ok" and day and day <= today:
            base = _parse_date(p["checked"])
            if base is None or day > base:
                p["checked_in_base"] = p["checked"]
                p["checked"] = day.isoformat()
                p["age_days"] = (today - day).days
                p["stale"] = p["age_days"] > STALE_DAYS
        elif v.get("status") == "changed":
            p["stale"] = True
            p["source_rate_pct"] = v.get("found_rate_pct")
    return programs


def _read_json(path: Path):
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {}
    return {}


def build(xlsx_path: Path, out_path: Path, today: dt.date, offline: bool = False):
    previous = _read_json(out_path).get("key_rate") or {}
    if offline:
        results, errors = [], []
    else:
        results, errors = collect_key_rate(today)
    key_rate, warnings = reconcile_key_rate(results, errors, previous, today)
    if offline and key_rate["value_pct"] is None:
        warnings.append("ключевая ставка недоступна")

    programs, prog_warnings = load_programs(xlsx_path, today)
    warnings += prog_warnings
    programs = apply_verification(programs, _read_json(out_path.parent / "verification.json"), today)
    for p in programs:
        if p.get("verification") == "changed":
            warnings.append(f"{p['bank']}, {p['name']}: на источнике другая ставка ({p.get('source_rate_pct')}%), обновите базу")

    rates = [p["rate_min_pct"] for p in programs if p["rate_min_pct"] is not None]
    best = min(rates) if rates else None
    spread = None
    if best is not None and key_rate["value_pct"] is not None:
        spread = round(best - key_rate["value_pct"], 2)

    result = {
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "key_rate": key_rate,
        "summary": {
            "programs": len(programs),
            "stale_programs": sum(1 for p in programs if p["stale"]),
            "lowest_rate_pct": best,
            "lowest_rate_spread_over_key_pct": spread,
        },
        "programs": programs,
        "warnings": warnings,
    }
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return result


def check(out_path: Path, today: dt.date):
    """Проверяет готовый market.json. Возвращает список проблем (пустой список = всё свежо)."""
    data = _read_json(out_path)
    kr = data.get("key_rate") or {}
    problems = []
    if kr.get("value_pct") is None:
        problems.append("ключевая ставка отсутствует")
    if kr.get("status") == "conflict":
        problems.append("источники ЦБ дали разные значения ставки")
    age = _age_days(kr.get("confirmed_at"), today)
    if age is None:
        problems.append("неизвестно, когда ставка подтверждалась")
    elif age > KEY_RATE_MAX_AGE_DAYS:
        problems.append(f"ставка не подтверждалась {age} дн. (допустимо {KEY_RATE_MAX_AGE_DAYS})")
    return problems


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    here = Path(__file__).resolve().parent
    ap.add_argument("--xlsx", type=Path, default=here / "baza-bankov.xlsx")
    ap.add_argument("--out", type=Path, default=here / "data" / "market.json")
    ap.add_argument("--offline", action="store_true", help="не обращаться к ЦБ (для проверки)")
    ap.add_argument("--check", action="store_true", help="проверить свежесть готового market.json")
    args = ap.parse_args(argv)

    if args.check:
        problems = check(args.out, dt.date.today())
        if problems:
            for p in problems:
                print("НЕ СВЕЖО:", p)
            return 1
        print("Ставка ЦБ подтверждена свежо.")
        return 0

    if not args.xlsx.exists():
        print(f"Не найден файл базы: {args.xlsx}", file=sys.stderr)
        return 1
    res = build(args.xlsx, args.out, dt.date.today(), args.offline)
    kr = res["key_rate"]
    print(f"Ключевая ставка: {kr['value_pct']}%, подтверждена {kr['confirmed_at']}, "
          f"изменена {kr['effective_since']}, источники: {', '.join(kr['sources']) or 'нет'} ({kr['status']})")
    if kr["last_meeting"]:
        print(f"Последнее заседание {kr['last_meeting']['date']}: ставка {kr['last_meeting']['decision']}")
    s = res["summary"]
    print(f"Программ: {s['programs']}, устарело: {s['stale_programs']}, минимальная ставка: {s['lowest_rate_pct']}%")
    for w in res["warnings"]:
        print("ВНИМАНИЕ:", w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
